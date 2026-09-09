"""文件格式转换：把 xlsx / pdf 转成 markdown 文本，供 LLM 阅读。

支持格式：.md / .txt（原文返回）、.pdf（pymupdf 提取）、.xlsx（openpyxl 提取表格）。
"""

from __future__ import annotations

from pathlib import Path

import pymupdf
from openpyxl import load_workbook

MAX_CELL_LEN = 200  # 单元格超长时截断，避免撑爆上下文


def xlsx_to_markdown(path: Path) -> str:
    """把 xlsx 每个工作表转成 markdown 表格。"""
    wb = load_workbook(path, read_only=True, data_only=True)
    parts: list[str] = []
    try:
        for ws in wb.worksheets:
            parts.append(f"## 工作表: {ws.title}\n")
            rows = [list(r) for r in ws.iter_rows(values_only=True)]
            rows = _clean_rows(rows)
            if not rows:
                parts.append("(空表)\n")
                continue
            parts.append(_rows_to_md_table(rows))
            parts.append("")
    finally:
        wb.close()
    return "\n".join(parts)


def _clean_rows(rows: list[list]) -> list[list]:
    """去掉全空行、截掉末尾的空列（合并单元格标题会产生大量空列）。"""
    rows = [r for r in rows if any(c not in (None, "") for c in r)]
    if not rows:
        return []
    width = max(max(i for i, c in enumerate(r) if c not in (None, "")) for r in rows) + 1
    return [r[:width] for r in rows]


def _cell(v) -> str:
    if v is None:
        return ""
    s = str(v).replace("\n", " ").replace("|", "\\|")
    return s if len(s) <= MAX_CELL_LEN else s[: MAX_CELL_LEN - 1] + "…"


def _rows_to_md_table(rows: list[list]) -> str:
    header = rows[0]
    lines = ["| " + " | ".join(_cell(c) for c in header) + " |"]
    lines.append("|" + "---|" * len(header))
    for row in rows[1:]:
        row = list(row) + [""] * (len(header) - len(row))  # 补齐短行
        lines.append("| " + " | ".join(_cell(c) for c in row[: len(header)]) + " |")
    return "\n".join(lines)


def pdf_to_markdown(path: Path) -> str:
    """按页提取 PDF 文本，每页用二级标题标注页码。"""
    parts: list[str] = []
    with pymupdf.open(path) as doc:
        for i, page in enumerate(doc, start=1):
            text = page.get_text("text").strip()
            parts.append(f"## 第 {i} 页\n")
            parts.append(text if text else "(无文本，可能是扫描图片页)")
            parts.append("")
    return "\n".join(parts)


def to_markdown(path: Path) -> str:
    """按扩展名分发到对应转换器。"""
    ext = path.suffix.lower()
    if ext in (".md", ".markdown", ".txt"):
        return path.read_text(encoding="utf-8", errors="replace")
    if ext == ".pdf":
        return pdf_to_markdown(path)
    if ext in (".xlsx", ".xlsm"):
        return xlsx_to_markdown(path)
    raise ValueError(f"不支持的文件格式: {ext}（目前支持 md/txt/pdf/xlsx）")
